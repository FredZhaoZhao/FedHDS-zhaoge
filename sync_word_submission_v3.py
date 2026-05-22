from __future__ import annotations

import copy
import shutil
import zipfile
from datetime import datetime
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm


ROOT = Path(__file__).resolve().parent
SOURCE_DOCX = ROOT / "paper_draft_submission_v2_uestc_template.docx"
TARGET_DOCX = ROOT / "paper_draft_submission_v3_uestc_template.docx"
FIG1_PATH = ROOT / "formal_cloud_results" / "stage2_dsample04_round20_20260503" / "figures" / "convergence_fl_vs_fedhds.png"
FIG2_PATH = ROOT / "formal_cloud_results" / "stage2_dsample04_round20_20260503" / "figures" / "global_loss_final.png"
FIG3_PATH = ROOT / "formal_cloud_results" / "stage2_dsample04_round20_20260503" / "figures" / "forget_loss_before_after.png"


UPDATED_PARAGRAPHS: dict[int, str] = {
    4: (
        "Federated unlearning seeks to reduce the influence of a target client without fully "
        "retraining the federated model. In language-model federated instruction tuning, the "
        "challenge is to weaken the forget signal while preserving retained-client utility. "
        "We study a guarded post-training unlearning method on top of FedHDS: the forget-client "
        "gradient defines what should be removed, retained-client curvature defines the "
        "utility-preserving geometry, and global norm clipping, stepped trust-region checking, "
        "direction probing, and a forget-loss guard control the update. In a five-seed main suite "
        "on Dolly with Qwen2-0.5B under Dirichlet alpha=0.5, forget-batch Hessian gives the "
        "strongest forgetting signal but is unstable, including one catastrophic utility failure "
        "and positive forgetting in only 4/5 valid seeds. By contrast, the retained-set "
        "negative-direction guarded update increases forget-client loss in 5/5 valid seeds while "
        "adding only +0.0618 mean global loss relative to the corresponding FedHDS baselines. GA "
        "and GA-guarded updates are stable but much weaker. Complementary scale checks with "
        "Qwen2.5-1.5B support a bounded conclusion: retained-set curvature becomes useful only "
        "after direction correction and guard control, and the method should be read as a "
        "utility-preserving approximate unlearning framework rather than as the strongest "
        "forgetting update in every regime."
    ),
    19: (
        "The evaluation is intentionally narrow and focused. A feasibility study verifies the "
        "end-to-end pipeline, a five-seed main-suite experiment under Dirichlet alpha=0.5 "
        "compares forget-batch and guarded retained-set updates, a direction-and-guard ablation "
        "explains the sign problem, and complementary scale checks test whether the interpretation "
        "survives outside the core setting. The goal is not to exhaust every regime, but to "
        "determine whether retained-set curvature can support a controllable post-training "
        "unlearning update."
    ),
    23: (
        "3. A five-seed main-suite evaluation plus complementary scale checks support a bounded "
        "claim: the method is better understood as a utility-preserving and auditable approximate "
        "unlearning framework than as a universally strongest forgetting method, and its "
        "forgetting gain remains scale- and seed-sensitive."
    ),
    95: (
        "Finally, the seed-robustness group repeats the main trade-off configuration across seeds "
        "42, 43, 44, 45, 47 under Dirichlet alpha=0.5. One additional run, seed46, diverged "
        "during baseline training by round 2 (round2_global_loss = 10.0 for FL, FedHDS, and "
        "retrain_oracle), so it is archived as an anomalous appendix case rather than pooled with "
        "valid method comparisons. The replacement run seed47 restores the five-seed bundle and "
        "lets us separate method instability from training-time divergence."
    ),
    97: (
        "This study compares seven settings: FL, the FedHDS baseline, retrain_oracle, "
        "forget_batch_hessian, retained_negative_guard, ga, and ga_guarded. Together, these "
        "settings isolate baseline training quality, exact client exclusion, the effect of the "
        "curvature reference, and the roles of sign correction and guard control under the same "
        "forgetting signal."
    ),
    98: (
        "The forget-batch Hessian setting serves as the direct second-order unlearning baseline "
        "because it uses the forget-client batch as both the gradient source and the Hessian "
        "reference. By contrast, retained_negative_guard keeps the same forget-client gradient but "
        "replaces the Hessian reference with retained-client data, applies negative-direction "
        "correction, and validates the update through global-loss and forget-loss guards. GA and "
        "GA-guarded act as first-order controls with and without stepped guard enforcement. "
        "retrain_oracle is kept as an exact client-exclusion reference rather than as a practical "
        "deployment method."
    ),
    100: (
        "Reported metrics include final global loss, forget-client loss before and after "
        "unlearning, update L2 norm, numbers of accepted and attempted guard steps, guard "
        "outcome, and unlearning time. Utility preservation is assessed by whether the "
        "post-unlearning global loss remains close to the FedHDS baseline pre-unlearning loss or "
        "stays below the specified global-loss guard. Directional forgetting effectiveness is "
        "assessed by whether the forget-client loss is prevented from decreasing and, ideally, is "
        "increased."
    ),
    101: (
        "For the direction-and-guard ablation, interpretation focuses on the interaction among "
        "update sign, forget-loss guard, forget-client loss change, and global-loss guard "
        "behavior rather than on another FL-versus-baseline comparison. The scale study then "
        "evaluates the same retained-set negative-direction update under larger model and data "
        "settings, while the seed-robustness study separates conclusions that remain stable across "
        "seeds from those that are sensitive to the local training trajectory. We still report the "
        "current loss-based MIA values, but only as supplementary indicators; in this setup they "
        "are too weak to serve as the primary acceptance criterion for method selection."
    ),
    109: (
        "The main trade-off setting increases the data sampling ratio from 0.2 to 0.4 and the "
        "number of training rounds from 10 to 20, and serves as the primary testbed for the "
        "utility-forgetting trade-off. The key quantitative claims in this section are based on "
        "the five valid seeds 42, 43, 44, 45, 47 under Dirichlet alpha=0.5."
    ),
    110: "Table 3. Five-seed aggregate utility-forgetting results under the main Dirichlet alpha=0.5 setting.",
    112: (
        "Figure 1. A representative FL vs FedHDS training convergence run (data sampling ratio "
        "0.4, 20 rounds). In this representative run, FedHDS converges to a lower final global "
        "loss than FL. Across the five valid seeds, however, the mean baseline gap is small, so "
        "the unlearning comparisons are anchored to the same within-seed FedHDS checkpoint rather "
        "than to a universal claim that FedHDS always dominates FL."
    ),
    113: (
        "For visual clarity, Figures 2 and 3 show representative runs from the same "
        "configuration, whereas the quantitative claims below rely on the five-seed aggregate in "
        "Table 3. Across these five valid seeds, FL and FedHDS remain close as training "
        "baselines, with a mean final-global-loss gap of only -0.0171 in favor of FL. The "
        "unlearning comparisons are therefore interpreted relative to the same pre-unlearning "
        "FedHDS checkpoint within each seed, which removes seed-level training variation from the "
        "utility comparison."
    ),
    114: "Figure 2. Final global loss comparison for a representative run in the main trade-off configuration.",
    115: (
        "The aggregate result sharpens the single-run story. forget_batch_hessian gives the "
        "largest mean forgetting increase (+3.6769), but it does so at a much larger mean utility "
        "cost (+3.7339 global loss relative to the same-seed FedHDS baseline) and with "
        "substantial instability. In the five valid seeds, it is positive on forgetting in only "
        "4/5 cases, and one seed (seed42) collapses to final global loss 18.5533. Another seed "
        "(seed44) even shows a slight negative forget delta (-0.0109). The direct forget-batch "
        "curvature reference is therefore strong but unreliable."
    ),
    116: (
        "Figure 3 shows the representative sign-misalignment case from the forget-client side, "
        "which motivated the later negative-direction correction and guard design."
    ),
    117: "Figure 3. Forget-client loss before and after unlearning for a representative run in the main trade-off configuration.",
    118: (
        "By contrast, retained_negative_guard is the most balanced second-order method in the "
        "five-seed suite. It increases forget-client loss in 5/5 valid seeds, with mean "
        "Deltaforget = +0.4011, while raising global loss by only +0.0618 on average relative to "
        "the same-seed FedHDS baselines. Its mean update norm is only 0.0636, and the stepped "
        "guard accepts 3.2/4.0 steps on average before stopping, usually because of the "
        "global-loss guard. This makes the method auditable and repeatable, even though the "
        "forgetting gain is smaller than that of the most aggressive forget-batch update. The two "
        "first-order controls, ga and ga_guarded, are even more utility-preserving, but their "
        "forgetting gains are much weaker. Both methods remain positive on forgetting in 5/5 "
        "seeds, yet their mean forget increases are only about +0.01, more than an order of "
        "magnitude smaller than that of retained_negative_guard. They are therefore useful "
        "conservative controls, but not strong standalone unlearning solutions in this setting. "
        "The retrain reference also clarifies the boundary of what this split can support. Across "
        "the five valid seeds, retrain_oracle leaves the tracked forget-loss proxy unchanged "
        "(Deltaforget = 0) and stays near the FedHDS utility level (Deltaglobal = +0.0012). It "
        "remains a useful exclusion reference, but not a stronger forgetting benchmark than the "
        "guarded post-training updates."
    ),
    125: (
        "Table 5 shows the correction result once the sign is flipped. The maximum-norm-0.05 case "
        "is the most conservative and gives the smallest global-loss increase. The maximum-norm-"
        "0.10 case provides the clearest forgetting gain while remaining within controlled utility "
        "degradation, and is therefore the preferred corrected setting. The maximum-norm-0.15 "
        "case is useful mainly as a guard-stopping example: the stepped guard truncates the "
        "update at 4/5 steps before the full requested norm is applied. Additional norm-sweep "
        "plots are omitted from the main text because they do not change this conclusion. The "
        "five-seed aggregate in Table 3 is consistent with this mechanism diagnosis: once the sign "
        "is corrected and the guards are enforced, retained_negative_guard remains positive on "
        "forgetting in all five valid seeds."
    ),
    128: (
        "The scale study is kept compact because its role is confirmatory rather than "
        "exploratory. Table 6 summarizes the 2 x 2 matrix over model size and data sampling ratio "
        "using one representative retained-set negative-direction update (maximum norm 0.10) and "
        "the forget-batch baseline."
    ),
    131: "5.5 Seed Robustness under the Main Dirichlet alpha=0.5 Setting",
    132: (
        "Table 3 already provides the main seed-robustness evidence under the central "
        "Qwen2-0.5B, data-sampling-ratio-0.4, 20-round setting. Across the five valid seeds 42, "
        "43, 44, 45, 47, retained_negative_guard is positive on forgetting in 5/5 cases, ga and "
        "ga_guarded are also positive in 5/5 cases but remain much weaker, and "
        "forget_batch_hessian is positive in only 4/5 cases while exhibiting much larger utility "
        "variance. The stable conclusion across seeds is therefore not that every method forgets "
        "equally well, but that sign-corrected retained-set curvature is the only non-trivial "
        "second-order update in this suite that remains positive on forgetting in every valid seed "
        "while keeping the utility cost bounded. An additional run, seed46, is not pooled into "
        "the main comparison because it fails earlier than the unlearning stage. Its baseline runs "
        "(FL, FedHDS, and retrain_oracle) all reach round2_global_loss = 10.0, which indicates "
        "training-time divergence rather than an unlearning-specific failure. The run is therefore "
        "archived as an appendix anomaly instead of being treated as method evidence. This "
        "distinction matters: the variability that remains after removing the anomalous seed is "
        "primarily about the magnitude of forgetting gain, not about whether guarded retained-set "
        "updates flip to the wrong sign."
    ),
    134: (
        "This study remains limited to Dolly, the Qwen/Qwen2.5 family, one federated "
        "configuration, and a single designated forget client. Forget-client loss is still an "
        "operational proxy rather than a gold-standard forgetting metric, and the five-seed suite "
        "shows that the magnitude of forgetting gain remains trajectory-sensitive even after an "
        "obviously divergent seed is removed. The current loss-based MIA values are also too weak "
        "to carry the main argument and are therefore treated only as supplementary indicators. "
        "Future work should test broader datasets, multi-client removal, stronger privacy audits, "
        "and cheaper retained-set HVP implementations."
    ),
    136: (
        "This paper studies post-training federated unlearning on top of a fixed FedHDS "
        "checkpoint. First, it establishes a protocol that isolates unlearning behavior from "
        "federated training quality and verifies that the full pipeline can be measured end to "
        "end. Second, the five-seed main suite shows that retained-set curvature becomes useful "
        "only after direction validation and guard control: the sign-corrected guarded update is "
        "positive on forgetting in all five valid seeds while adding only +0.0618 mean global "
        "loss relative to the corresponding FedHDS baselines. Third, the same suite also "
        "clarifies the boundary of the method: forget-batch curvature can forget more "
        "aggressively, but it is substantially less stable, while GA-style guarded first-order "
        "updates are stable but too weak to serve as the main unlearning mechanism. The resulting "
        "claim is therefore deliberately bounded. The proposed method is best understood as a "
        "utility-preserving and auditable approximate unlearning framework, not as a universally "
        "strongest forgetting update in every regime."
    ),
}


TABLE3_ROWS = [
    [
        "Method",
        "Mean Deltaforget",
        "Positive seeds",
        "Mean Deltaglobal vs FedHDS",
        "Mean update L2",
        "Mean accepted steps",
        "Mean runtime (s)",
        "Main reading",
    ],
    ["FedHDS baseline", "N/A", "N/A", "+0.0000", "N/A", "N/A", "N/A", "reference checkpoint"],
    ["FL baseline", "N/A", "N/A", "-0.0171", "N/A", "N/A", "N/A", "comparable training baseline"],
    ["Retrain baseline", "+0.0000", "0/5", "+0.0012", "N/A", "0/0", "0.00", "exact exclusion reference, not stronger forgetting"],
    ["Forget-batch Hessian", "+3.6769", "4/5", "+3.7339", "2.4336", "1/1", "1.72", "strongest forgetting, unstable utility"],
    ["Retained-set negative guard", "+0.4011", "5/5", "+0.0618", "0.0636", "3.2/4.0", "112.30", "main utility-preserving trade-off"],
    ["GA", "+0.0096", "5/5", "+0.0016", "0.0055", "1/1", "1.26", "stable but weak forgetting"],
    ["GA guarded", "+0.0097", "5/5", "-0.0011", "0.0055", "5/5", "125.45", "most utility-stable, forgetting still weak"],
]


def replace_paragraph_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run.text = ""


def update_table_3(table) -> None:
    current_rows = len(table.rows)
    current_cols = len(table.rows[0].cells)

    while current_rows < len(TABLE3_ROWS):
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))
        current_rows += 1

    if current_cols != len(TABLE3_ROWS[0]):
        raise RuntimeError(f"Table 3 column mismatch: expected {len(TABLE3_ROWS[0])}, found {current_cols}")

    for r_idx, row_values in enumerate(TABLE3_ROWS):
        row = table.rows[r_idx]
        for c_idx, value in enumerate(row_values):
            row.cells[c_idx].text = value

    # Clear any leftover rows if the source table had more rows than needed.
    for r_idx in range(len(TABLE3_ROWS), len(table.rows)):
        for cell in table.rows[r_idx].cells:
            cell.text = ""


def find_paragraph_index(doc: Document, startswith: str) -> int:
    for idx, para in enumerate(doc.paragraphs):
        if para.text.strip().startswith(startswith):
            return idx
    raise RuntimeError(f"Could not find paragraph starting with: {startswith}")


def insert_picture_after(paragraph, image_path: Path, width_cm: float = 14.5) -> None:
    pic_para = paragraph._parent.add_paragraph()
    pic_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = pic_para.add_run()
    run.add_picture(str(image_path), width=Cm(width_cm))
    paragraph._p.addnext(pic_para._p)


def count_docx_assets(path: Path) -> tuple[int, int, int]:
    with zipfile.ZipFile(path, "r") as zf:
        names = zf.namelist()
        media_count = len(
            [name for name in names if name.startswith("word/media/") and not name.endswith("/")]
        )
        xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")
    omath_count = len(re.findall(r"<m:oMath", xml))
    drawing_count = len(re.findall(r"<w:drawing", xml))
    return media_count, omath_count, drawing_count


def main() -> None:
    if not SOURCE_DOCX.exists():
        raise FileNotFoundError(SOURCE_DOCX)

    temp_target = TARGET_DOCX.with_name(f"{TARGET_DOCX.stem}.tmp.docx")
    if temp_target.exists():
        temp_target.unlink()

    if TARGET_DOCX.exists():
        backup_current = TARGET_DOCX.with_name(
            f"{TARGET_DOCX.stem}.before_wps_rebuild_{datetime.now().strftime('%Y%m%d_%H%M%S')}.bak.docx"
        )
        shutil.copy2(TARGET_DOCX, backup_current)

    shutil.copy2(SOURCE_DOCX, temp_target)

    doc = Document(temp_target)

    for idx, text in UPDATED_PARAGRAPHS.items():
        replace_paragraph_text(doc.paragraphs[idx], text)

    update_table_3(doc.tables[2])

    fig1_anchor = doc.paragraphs[find_paragraph_index(doc, "Figure 1 plots the per-round convergence of FL and FedHDS.")]
    fig2_anchor = doc.paragraphs[find_paragraph_index(doc, "For visual clarity, Figures 2 and 3 show representative runs")]
    fig3_anchor = doc.paragraphs[find_paragraph_index(doc, "Figure 3 shows the representative sign-misalignment case")]
    insert_picture_after(fig1_anchor, FIG1_PATH)
    insert_picture_after(fig2_anchor, FIG2_PATH)
    insert_picture_after(fig3_anchor, FIG3_PATH)

    doc.save(temp_target)

    media_count, omath_count, drawing_count = count_docx_assets(temp_target)
    if media_count < 28 or omath_count < 18 or drawing_count < 27:
        raise RuntimeError(
            f"Asset preservation check failed: media={media_count}, omath={omath_count}, drawing={drawing_count}"
        )

    shutil.move(str(temp_target), str(TARGET_DOCX))

    media_count, omath_count, drawing_count = count_docx_assets(TARGET_DOCX)
    print(f"Updated {TARGET_DOCX.name}")
    print(f"media={media_count} omath={omath_count} drawing={drawing_count}")


if __name__ == "__main__":
    main()
