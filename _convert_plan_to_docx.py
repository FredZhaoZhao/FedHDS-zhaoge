# -*- coding: utf-8 -*-
"""Convert 第三次补强计划.md to 第三次补强计划.docx with proper formatting."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


MD_PATH = Path(r"F:\FedHDS-zhaoge\第三次补强计划.md")
DOCX_PATH = Path(r"F:\FedHDS-zhaoge\第三次补强计划.docx")


def set_run_font(run, name="Times New Roman", cjk="SimSun", size=10.5, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), cjk)


def add_heading_para(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    sizes = {1: 18, 2: 15, 3: 13, 4: 11}
    size = sizes.get(level, 11)
    run = p.add_run(text)
    set_run_font(run, name="Times New Roman", cjk="SimHei", size=size, bold=True)
    return p


def add_body_para(doc, text, bullet=False, indent=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.4
    if indent:
        p.paragraph_format.left_indent = Cm(0.7 * indent)
    if bullet:
        text = "• " + text
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, name="Consolas", cjk="SimSun", size=10)
        else:
            run = p.add_run(part)
            set_run_font(run)
    return p


def add_code_block(doc, code_lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(2)
    for i, line in enumerate(code_lines):
        if i > 0:
            p.add_run().add_break()
        run = p.add_run(line)
        set_run_font(run, name="Consolas", cjk="SimSun", size=9.5)


def add_table_block(doc, rows):
    if not rows:
        return
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Light Grid Accent 1"
    for i, row_cells in enumerate(rows):
        for j, cell_text in enumerate(row_cells):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(cell_text.strip())
            set_run_font(run, size=9.5, bold=(i == 0))
    doc.add_paragraph()


def parse_table(lines, idx):
    table_lines = []
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        table_lines.append(lines[idx])
        idx += 1
    rows = []
    for i, ln in enumerate(table_lines):
        if i == 1 and re.match(r"^\s*\|[\s\-:|]+\|\s*$", ln):
            continue
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    return rows, idx


def convert():
    text = MD_PATH.read_text(encoding="utf-8")
    lines = text.split("\n")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    rPr = style.element.get_or_add_rPr()
    rFonts = OxmlElement("w:rFonts")
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "SimSun")
    rPr.append(rFonts)

    i = 0
    in_code = False
    code_buffer = []
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(doc, code_buffer)
                code_buffer = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        if stripped.startswith("|") and "|" in stripped[1:]:
            rows, i = parse_table(lines, i)
            add_table_block(doc, rows)
            continue

        if not stripped:
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            add_heading_para(doc, m.group(2), level)
            i += 1
            continue

        if stripped == "---":
            p = doc.add_paragraph()
            pPr = p._element.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), "808080")
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        m = re.match(r"^(\s*)[-*]\s+\[[ x]\]\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            checked = "[x]" in line.lower()
            mark = "☑ " if checked else "☐ "
            add_body_para(doc, mark + m.group(2), bullet=False, indent=indent + 1)
            i += 1
            continue

        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            add_body_para(doc, m.group(2), bullet=True, indent=indent)
            i += 1
            continue

        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            add_body_para(doc, f"{m.group(2)}. {m.group(3)}", indent=indent)
            i += 1
            continue

        add_body_para(doc, stripped)
        i += 1

    doc.save(DOCX_PATH)
    print(f"Saved: {DOCX_PATH}")


if __name__ == "__main__":
    convert()
